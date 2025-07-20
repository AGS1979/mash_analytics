import os
import re
from typing import List, Tuple

import pdfplumber
import requests
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.shared import RGBColor
import yfinance as yf


# ==========================
# DeepSeek Setup
# ==========================
DEEPSEEK_API_KEY = os.getenv("DEEPSEEK_API_KEY")  # Ensure it's set in your .env or environment
DEEPSEEK_URL = "https://api.deepseek.com/v1/chat/completions"
FMP_API_KEY = os.getenv("FMP_API_KEY")

# ==========================
# Report Structures
# ==========================
REPORT_TEMPLATES = {
    "Spin-Off or Split-Up": """
Transaction Overview
ParentCo and SpinCo details
Rationale (regulatory, strategic unlock, valuation arbitrage)
Distribution terms (ratio, eligibility, tax treatment)
ParentCo Post-Spin Outlook
Strategic focus
Financial profile and valuation
SpinCo Investment Case
Business model, growth drivers
Historical and pro forma financials
Independent valuation (e.g., Sum-of-the-Parts)
Valuation Analysis
Risks and Overhangs
Forced selling, low float, governance concerns
""",
    "Mergers & Acquisitions": """
Deal Summary
Parties involved, consideration (cash/stock), premium
Regulatory/antitrust/board approval status
Target Company Analysis
Valuation vs. offer
Control premium vs. peers
Buyer’s Rationale and Financing
Strategic fit
Synergies and pro forma financials
Deal financing (debt, equity)
Shareholder Vote & Antitrust Risk
Key holders' stance
Timing and likelihood of deal closure
Spread Analysis and Arbitrage Opportunity
Deal spread
IRR scenarios based on timing/risk
""",
    "Bankruptcy / Distressed / Restructuring": """
Situation Summary
Cause of distress
Filing date, jurisdiction, DIP terms
Capital Structure Analysis
Pre- and post-reorg structure
Seniority waterfall
Creditor classes and recovery potential
Valuation and Recovery Scenarios
Estimated Enterprise Value
Recovery per instrument (bonds, equity, unsecured)
Reorganization Plan and Exit Timeline
Conversion to equity, rights offering, warrants
Exit multiples
Catalysts and Legal Risks
Judge approval, creditor objections, asset sales
""",
    "Activist Campaign": """
Activist Background
Fund profile, history, prior campaigns
Campaign Details
Demands (board seat, spin, buyback, etc.)
Timeline of engagement
Company's Response and Governance Profile
Management alignment, shareholder defense
Scenario Analysis
Status quo vs. activist success
Proxy fight implications
Valuation Impact
NPV of potential changes (e.g., spin-off value, ROIC uplift)
""",
    "Regulatory or Legal Catalyst": """
Legal/Regulatory Background
Case/issue summary
Historical legal proceedings
Outcome Scenarios
Win, loss, settlement
Timeline
Financial and Strategic Implications
Fines, product approval, license loss
Revenue/EBITDA impact
Market Reaction History (if any)
Past similar cases
""",
    "Asset Sales or Carve-Outs": """
Transaction Overview
Buyer, price, structure
Valuation vs. book and peers
Strategic Impact
Focus shift, deleveraging, margin profile
Use of Proceeds
Debt repayment, dividends, buybacks, capex
Re-rating Potential
EBITDA margin uplift, return metrics
""",
    "Capital Raising or Buyback Catalyst": """
Transaction Mechanics
Size, dilution, instrument type
Capital Structure Post-Deal
Leverage ratios, interest burden
Shareholder Implications
Accretion/dilution
EPS impact
Buyback Analysis (if applicable)
Repurchase pace, valuation support
"""
}


# ==========================
# Financial Data Helpers
# ==========================

def resolve_company_to_ticker(company_name: str) -> str:
    prompt = f"What is the stock ticker for the public company '{company_name}'?"
    headers = {"Authorization": f"Bearer {DEEPSEEK_API_KEY}"}
    payload = {
        "model": "deepseek-chat",
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0
    }
    try:
        res = requests.post(DEEPSEEK_URL, headers=headers, json=payload)
        res.raise_for_status()
        ticker = res.json()["choices"][0]["message"]["content"].strip()
        # Sanitize to get a clean ticker format
        return re.sub(r'[^A-Z\.]', '', ticker)
    except Exception:
        return None

def get_ev_ebitda_multiple(ticker: str) -> float:
    if not FMP_API_KEY:
        print("FMP_API_KEY not set. Skipping EV/EBITDA fetch.")
        return 0.0
    url = f"https://financialmodelingprep.com/api/v3/key-metrics-ttm/{ticker}?apikey={FMP_API_KEY}"
    try:
        r = requests.get(url)
        r.raise_for_status()
        data = r.json()
        if isinstance(data, list) and data:
            multiple = data[0].get("enterpriseValueOverEBITDATTM")
            return float(multiple) if multiple is not None else 0.0
    except Exception:
        return 0.0
    return 0.0

def fetch_fundamentals_yf(ticker: str) -> Tuple[float, float, float]:
    """Returns (market_cap, net_debt, ttm_ebitda) via Yahoo Finance."""
    try:
        t = yf.Ticker(ticker)
        info = t.info or {}
        market_cap = info.get("marketCap", 0) or 0
        total_debt = info.get("totalDebt", 0) or 0
        cash = info.get("cashAndShortTermInvestments", info.get("cash", 0)) or 0
        net_debt = total_debt - cash
        ebitda = info.get("ebitda", 0) or 0
        return float(market_cap), float(net_debt), float(ebitda)
    except Exception:
        return 0.0, 0.0, 0.0


# ==========================
# Text Extractors
# ==========================

def extract_text_from_pdf(file_path: str) -> str:
    try:
        text = ""
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        return text.strip()
    except Exception as e:
        return f"[ERROR extracting PDF: {e}]"

def extract_text_from_docx(file_path: str) -> str:
    try:
        doc = Document(file_path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as e:
        return f"[ERROR extracting DOCX: {e}]"

# ==========================
# Main Orchestration Function
# ==========================

def clean_markdown(text):
    text = re.sub(r'#+\s*', '', text)
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    text = re.sub(r'`{1,3}(.*?)`{1,3}', r'\1', text)
    text = re.sub(r'!\[.*?\]\(.*?\)', '', text)
    text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    text = re.sub(r'^- ', '• ', text, flags=re.MULTILINE)
    return text.strip()

def truncate_safely(text, limit=7000):
    if len(text) <= limit:
        return text
    cutoff = text[:limit].rfind('\n\n')
    return text[:cutoff] if cutoff != -1 else text[:limit]

def generate_special_situation_note(
    company_name: str,
    situation_type: str,
    file_paths: List[str],
    output_path: str,
    valuation_mode: str = None,
    parent_peers: str = "",
    spinco_peers: str = ""
):
    # 1) Build the combined_text from all inputs
    combined_text = ""
    # ✨ START: CORRECTED CODE BLOCK
    for path in file_paths:
        if path.lower().endswith(".pdf"):
            combined_text += extract_text_from_pdf(path) + "\n\n"
        elif path.lower().endswith(".docx"):
            combined_text += extract_text_from_docx(path) + "\n\n"
        else:
            combined_text += f"[Unsupported file: {path}]\n\n"
    # ✨ END: CORRECTED CODE BLOCK

    # 2) Grab the template structure
    structure = REPORT_TEMPLATES.get(situation_type)
    if not structure:
        raise ValueError(f"Unsupported situation type: {situation_type}")

    # 3) Build the valuation section based on user input
    valuation_section = ""
    if situation_type == "Spin-Off or Split-Up" and valuation_mode:
        
        def process_peers(raw_peers_str: str):
            names = [n.strip() for n in raw_peers_str.split(',') if n.strip()]
            tickers = [resolve_company_to_ticker(n) for n in names]
            valid_tickers = [t for t in tickers if t]
            multiples = [get_ev_ebitda_multiple(t) for t in valid_tickers]
            valid_multiples = [m for m in multiples if m and m > 0]
            avg_multiple = round(sum(valid_multiples) / len(valid_multiples), 2) if valid_multiples else None
            return names, valid_multiples, avg_multiple

        if valuation_mode == "ai_peers":
            valuation_section = "For the Valuation Analysis section, please identify relevant public peer companies for the ParentCo and SpinCo. Use their average LTM EV/EBITDA multiples to perform a Sum-of-the-Parts (SOTP) valuation based on the TTM EBITDA figures found in the provided documents. Compare the resulting implied equity value to the parent company's current market capitalization to estimate the potential value unlock."

        elif valuation_mode == "user_peers":
            p_names, p_mults, p_avg = process_peers(parent_peers)
            s_names, s_mults, s_avg = process_peers(spinco_peers)
            valuation_section = f"""
For the Valuation Analysis section, use the following user-provided peer data:

**ParentCo Peers**: {', '.join(p_names)}
- EV/EBITDA multiples: {p_mults} (Average: {p_avg or 'N/A'})

**SpinCo Peers**: {', '.join(s_names)}
- EV/EBITDA multiples: {s_mults} (Average: {s_avg or 'N/A'})

Apply these average multiples to the respective TTM EBITDA figures from the documents to perform a Sum-of-the-Parts (SOTP) valuation.
"""

    # 4) Build the enhanced prompt
    prompt = f"""
You are an institutional investment analyst writing a professional memo on a special situation involving {company_name}.
The situation is: **{situation_type}**

Below is the internal company information extracted from various files:
\"\"\"{truncate_safely(combined_text)}\"\"\"

{valuation_section}

Using the structure below, generate a detailed, data-driven investment memo.
Structure:
{structure}
"""

    # 5) Call DeepSeek and format the output (Note: Corrected step numbering from 4 to 5)
    headers = {"Authorization": f"Bearer {DEEPSEEK_API_KEY}"}
    payload = {
        "model": "deepseek-chat",
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0.3
    }
    response = requests.post(DEEPSEEK_URL, headers=headers, json=payload)
    response.raise_for_status()
    memo = clean_markdown(response.json()["choices"][0]["message"]["content"])


    memo_dict = split_into_sections(memo, structure)
    format_memo_docx(memo_dict, company_name, situation_type, output_path)


def format_memo_docx(memo_dict, company_name, situation_type, output_path):
    doc = Document()

    # 2) Set default Normal style to 12 pt, black, Aptos Display
    normal = doc.styles['Normal']
    normal.font.name = 'Aptos Display'
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    normal.paragraph_format.line_spacing = 1.15

    # 3) Title (centered, 24 pt)
    title = doc.add_heading(level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"{company_name} – {situation_type} Investment Memo")
    run.font.name = 'Aptos Display'
    run.font.size = Pt(24)
    run.bold = True

    doc.add_paragraph()  # spacer

    if not isinstance(memo_dict, dict) or not memo_dict:
        memo_dict = {"Executive Summary": str(memo_dict)}

    for section_title, content in memo_dict.items():
        # 4) Use built-in Heading 1 style for section titles
        heading = doc.add_heading(section_title, level=1)
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.space_after = Pt(6)

        # Bold the first sentence of each section as a mini-lead
        paras = content.strip().split('\n\n')
        for idx, para in enumerate(paras):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.2

            # first sentence bolded
            sentences = re.split(r'(?<=[.?!]) +', para.strip())
            if sentences:
                run0 = p.add_run(sentences[0].strip() + " ")
                run0.bold = True
                run0.font.size = Pt(12)
                run0.font.name = 'Aptos Display'
                run_rest = p.add_run(" ".join(sentences[1:]))
                run_rest.font.size = Pt(12)
                run_rest.font.name = 'Aptos Display'
            else:
                p.add_run(para.strip())

        # 5) Insert a blank paragraph for spacing
        doc.add_paragraph()

    # 6) Tighten page margins
    sec = doc.sections[0]
    for margin in ('left', 'right', 'top', 'bottom'):
        setattr(sec, f"{margin}_margin", Inches(0.75))

    doc.save(output_path)

def split_into_sections(text: str, template: str):
    """
    Splits the memo text into a dictionary of sections based on titles from a template.

    This corrected function reliably extracts section titles from the template and uses
    them as delimiters to split the memo text, avoiding the fragility of the previous
    generic pattern matching.

    Args:
        text: The memo text to be split.
        template: The report template string containing the section titles.

    Returns:
        A dictionary with section titles as keys and their content as values.
    """
    sections = {}
    # Extract canonical titles from the template, taking only text before any '('.
    titles = [line.split('(')[0].strip() for line in template.strip().split('\n') if line.strip()]
    if not titles:
        return {"Memo": text.strip()} if text.strip() else {}

    # Build a regex to find any of the titles when they appear on their own line.
    pattern = re.compile(r'^(' + '|'.join(map(re.escape, titles)) + r')\s*$', re.MULTILINE | re.IGNORECASE)

    matches = list(pattern.finditer(text))
    if not matches:
        return {"Memo": text.strip()} if text.strip() else {}

    # Iterate through the found titles to carve out the sections
    for i, match in enumerate(matches):
        title = match.group(1).strip()
        start_of_content = match.end()
        end_of_content = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        content = text[start_of_content:end_of_content].strip()
        
        # Use the canonical title from the template for consistent key names
        canonical_title = next((t for t in titles if t.lower() == title.lower()), title)
        sections[canonical_title] = content

    return sections