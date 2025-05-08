import os
import requests
from openai import OpenAI
import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
import pandas as pd

# --- Load secrets from environment ---
FMP_API_KEY = os.environ.get("FMP_API_KEY")
OPENAI_API_KEY = os.environ.get("OPENAI_API_KEY")
DEEPSEEK_API_KEY = os.environ.get("DEEPSEEK_API_KEY")

client = OpenAI(api_key=OPENAI_API_KEY)

# ─── Used by Flask to extract ticker ─────────────────────────────────────────
def process_query_1(query):
    prompt = f"Extract the stock ticker from this query using FMP tickers: '{query}'. Return ONLY the ticker."
    response = requests.post(
        "https://api.deepseek.com/chat/completions",
        headers={"Authorization": f"Bearer {DEEPSEEK_API_KEY}"},
        json={
            "model": "deepseek-chat",
            "messages": [{"role": "user", "content": prompt}],
            "temperature": 0,
        },
    )
    result = response.json()
    return result["choices"][0]["message"]["content"].strip().upper()


def get_fmp_json(endpoint):
    url = f"https://financialmodelingprep.com/api/v3/{endpoint}&apikey={FMP_API_KEY}"
    r = requests.get(url)
    return r.json() if r.ok else None


def add_dataframe_to_sheet(wb, sheet_name, df):
    ws = wb.create_sheet(sheet_name)
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)
    for col in ws.columns:
        for cell in col:
            cell.alignment = Alignment(wrap_text=True)
            if cell.row == 1:
                cell.font = Font(bold=True)


def summarize_with_openai(text, topic):
    prompt = f"Summarize the following {topic} in bullet points:\n{text}"
    response = client.chat.completions.create(
        model="gpt-4",
        messages=[{"role": "user", "content": prompt}],
        temperature=0.4
    )
    return response.choices[0].message.content


# ─── Called from Flask to create the Excel report ───────────────────────────
def create_stock_report(ticker):
    profile = get_fmp_json(f"profile/{ticker}?")
    if not profile:
        return None, None, None
    overview = profile[0]

    financials = pd.DataFrame(get_fmp_json(f"income-statement/{ticker}?limit=5"))
    estimates  = pd.DataFrame(get_fmp_json(f"analyst-estimates/{ticker}?limit=4"))
    ratings    = pd.DataFrame(get_fmp_json(f"rating/{ticker}?limit=4"))
    insider    = pd.DataFrame(get_fmp_json(f"insider-trading?symbol={ticker}&limit=10"))
    news       = pd.DataFrame(get_fmp_json(f"stock_news?tickers={ticker}&limit=10"))

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Overview sheet
    overview_ws = wb.create_sheet("Company Overview")
    overview_data = [
        ("Company Name", overview.get("companyName")),
        ("Ticker", ticker),
        ("Sector", overview.get("sector")),
        ("Industry", overview.get("industry")),
        ("Exchange", overview.get("exchange")),
        ("Market Cap", overview.get("mktCap")),
        ("Website", overview.get("website")),
        ("Description", overview.get("description")),
    ]
    for row in overview_data:
        overview_ws.append(row)

    # Add dataframes
    if not financials.empty:
        add_dataframe_to_sheet(wb, "Financials", financials)
    if not estimates.empty:
        add_dataframe_to_sheet(wb, "Estimates", estimates)
    if not ratings.empty:
        add_dataframe_to_sheet(wb, "Consensus", ratings)
    if not insider.empty:
        add_dataframe_to_sheet(wb, "Insider Trades", insider)
    if not news.empty:
        add_dataframe_to_sheet(wb, "News", news[["publishedDate", "title", "site", "url"]])

    # Investment insights
    investment_text = f"""
    Company: {overview.get("companyName")}
    Sector: {overview.get("sector")}
    Industry: {overview.get("industry")}
    Description: {overview.get("description")}
    """
    positives = summarize_with_openai(investment_text, "investment positives")
    risks     = summarize_with_openai(investment_text, "investment risks")

    pos_ws = wb.create_sheet("Investment Positives")
    pos_ws.append(["Key Positives"])
    for line in positives.split('\n'):
        pos_ws.append([line.strip()])

    risk_ws = wb.create_sheet("Investment Risks")
    risk_ws.append(["Key Risks"])
    for line in risks.split('\n'):
        risk_ws.append([line.strip()])

    # Save Excel file
    filename = f"{ticker}_Stock_Report.xlsx"
    full_path = os.path.join("Stock Reports", filename)
    os.makedirs("Stock Reports", exist_ok=True)
    wb.save(full_path)

    return full_path, wb, wb["Company Overview"]
